"""Repair a narrowly defined PowerPoint text-shape marker in a temporary copy."""

from __future__ import annotations

import io
import zipfile
from collections import defaultdict

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS = {"p": _PRESENTATION_NS}


def repair_unrecognized_text_shapes(content: bytes) -> bytes | None:
    """Mark unclassified nonempty text shapes as text boxes, preserving the source.

    Some presentation writers omit ``txBox`` on a ``p:sp`` containing text and
    no recognized geometry. python-pptx raises NotImplementedError for this
    shape, although its text is readable. Only repair that exact case. Any
    unclassified non-text shape or ambiguous XML leaves the document unchanged
    and lets the original parser error reach the user.
    """
    presentation = Presentation(io.BytesIO(content))
    targets: dict[str, set[int]] = defaultdict(set)

    def visit(shapes, part_name: str) -> bool:
        for shape in shapes:
            try:
                shape_type = shape.shape_type
            except NotImplementedError as exc:
                if str(exc) != "Shape instance of unrecognized shape type":
                    return False
                if not shape.has_text_frame or not shape.text.strip():
                    return False
                if shape._element.tag != f"{{{_PRESENTATION_NS}}}sp":
                    return False
                targets[part_name].add(shape.shape_id)
                continue
            if shape_type == MSO_SHAPE_TYPE.GROUP and not visit(
                shape.shapes, part_name
            ):
                return False
        return True

    for slide in presentation.slides:
        if not visit(slide.shapes, str(slide.part.partname).lstrip("/")):
            return None
    if not targets:
        return None

    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        names = source.namelist()
        if any(names.count(name) != 1 for name in targets):
            return None
        for info in source.infolist():
            data = source.read(info)
            if info.filename in targets:
                root = etree.fromstring(
                    data, etree.XMLParser(resolve_entities=False, no_network=True)
                )
                for shape_id in targets[info.filename]:
                    matches = root.xpath(
                        ".//p:sp[p:nvSpPr/p:cNvPr/@id=$shape_id]/p:nvSpPr/p:cNvSpPr",
                        namespaces=_NS,
                        shape_id=str(shape_id),
                    )
                    if len(matches) != 1:
                        return None
                    matches[0].set("txBox", "1")
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(info, data)
    return output.getvalue()
