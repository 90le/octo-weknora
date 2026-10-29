import io
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from docreader.parser.markitdown_parser import StdMarkitdownParser
from docreader.parser.pptx_text_shape_repair import repair_unrecognized_text_shapes


def presentation_bytes(text: str, *, omit_marker: bool) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    shape.text = text
    if omit_marker:
        shape._element.nvSpPr.cNvSpPr.attrib.pop("txBox", None)
        geometry = shape._element.spPr.find(qn("a:prstGeom"))
        if geometry is not None:
            shape._element.spPr.remove(geometry)
    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(1), Inches(3))
    result = io.BytesIO()
    presentation.save(result)
    return result.getvalue()


class PPTXTextShapeRepairTest(unittest.TestCase):
    def test_repair_preserves_text_and_other_package_parts(self):
        original = presentation_bytes("training marker\nsecond paragraph", omit_marker=True)
        with self.assertRaises(NotImplementedError):
            _ = Presentation(io.BytesIO(original)).slides[0].shapes[0].shape_type
        repaired = repair_unrecognized_text_shapes(original)
        self.assertIsNotNone(repaired)
        with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(
            io.BytesIO(repaired)
        ) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != "ppt/slides/slide1.xml":
                    self.assertEqual(before.read(name), after.read(name), name)
        shape = Presentation(io.BytesIO(repaired)).slides[0].shapes[0]
        self.assertEqual(shape.text, "training marker\nsecond paragraph")
        result = StdMarkitdownParser(
            file_name="training.pptx", file_type="pptx"
        ).parse_into_text(original)
        self.assertIn("training marker", result.content)
        self.assertIn("second paragraph", result.content)

    def test_does_not_rewrite_normal_or_unclassified_empty_shapes(self):
        self.assertIsNone(
            repair_unrecognized_text_shapes(presentation_bytes("normal", omit_marker=False))
        )
        self.assertIsNone(
            repair_unrecognized_text_shapes(presentation_bytes("", omit_marker=True))
        )
